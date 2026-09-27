import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="EFF6FF", border_hex="1E3A8A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(30, 58, 138)
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    
    # empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

doc = docx.Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Title
p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(2)
r_title = p_title.add_run("TUTORIEL MÉTHODOLOGIQUE PAS-À-PAS")
r_title.bold = True
r_title.font.name = "Calibri"
r_title.font.size = Pt(18)
r_title.font.color.rgb = RGBColor(30, 58, 138)

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(12)
r_sub = p_sub.add_run("Exercice 04 : Concevoir ses propres outils numériques de suivi sous Excel\nSpécialisation Préparateur Physique — HECh")
r_sub.font.name = "Calibri"
r_sub.font.size = Pt(12)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(100, 116, 139)

add_callout(doc, "Objectif de l'Exercice", 
    "Cet exercice a pour vocation de vous faire passer du statut de simple utilisateur d'outils à celui de concepteur de vos propres solutions de collecte et d'analyse. "
    "Vous construirez 5 outils modulaires sous Microsoft Excel répondant aux besoins quotidiens du staff : suivi de récupération (Hooper), charge interne (Foster RPE), carnet multi-sports, suivi nutritionnel et suivi longitudinal de la performance athlétique.")

# Table of contents summary
p_toc = doc.add_paragraph()
p_toc.paragraph_format.space_before = Pt(6)
p_toc.paragraph_format.space_after = Pt(6)
r_toc = p_toc.add_run("Sommaire des Étapes de Réalisation :")
r_toc.bold = True
r_toc.font.size = Pt(12)

steps_summary = [
    "Étape 1 : Conception de l'Outil 1 — Suivi de la Récupération (Indice de Hooper)",
    "Étape 2 : Conception de l'Outil 2 — Quantification de la Charge Interne (Foster RPE)",
    "Étape 3 : Conception de l'Outil 3 — Journal Numérique d'Entraînement Multi-Paramètres",
    "Étape 4 : Conception de l'Outil 4 — Suivi Nutritionnel & Hydratation Automatisé",
    "Étape 5 : Conception de l'Outil 5 — Batterie de Tests Physiques & Profil Athlétique",
    "Étape 6 : Synthèse Staff — Tableau de Bord Décisionnel et Alertes Prioritaires"
]
for s in steps_summary:
    p_s = doc.add_paragraph(s, style='List Bullet')
    p_s.paragraph_format.space_before = Pt(1)
    p_s.paragraph_format.space_after = Pt(2)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# ---------------------------------------------------------------------------
# ÉTAPE 1
# ---------------------------------------------------------------------------
h1 = doc.add_heading("Étape 1 : Outil 1 — Suivi de Récupération (Indice de Hooper)", level=1)
h1.paragraph_format.space_before = Pt(14)
h1.paragraph_format.space_after = Pt(6)

p1 = doc.add_paragraph(
    "L'indice de Hooper (1995) évalue le bien-être et la récupération subjective de l'athlète à travers 4 items cotés de 1 (très bon / nulle) à 7 (très mauvais / extrême) : "
    "le Sommeil, le Stress, la Fatigue et les Courbatures musculaires. Le score total oscille entre 4 (récupération parfaite) et 28 (épuisement sévère)."
)

add_callout(doc, "Formules et Paramétrage dans la feuille '01_Suivi_Hooper'",
    "1. Créez les en-têtes : Date (A), Athlète (B), Sommeil (C), Stress (D), Fatigue (E), Courbatures (F), Indice Hooper (G), Statut (H), Recommandation (I).\n"
    "2. Formule du Score Hooper en G5 : =SOMME(C5:F5)\n"
    "3. Formule du Statut en H5 : =SI(G5>=20; \"ALERTE FATIGUE\"; SI(G5>=15; \"Vigilance\"; \"Optimal\"))\n"
    "4. Formule de Recommandation Staff en I5 : =SI(G5>=20; \"Allègement 50% ou repos actif\"; SI(G5>=15; \"Adapter intensité / hydratation\"; \"Séance normale\"))\n"
    "5. Mise en forme conditionnelle : Appliquez un fond Rouge clair (#FEE2E2) si la valeur de la cellule G5 est >= 20, Orange (#FEF3C7) si >= 15, et Vert (#D1FAE5) sinon.")

# ---------------------------------------------------------------------------
# ÉTAPE 2
# ---------------------------------------------------------------------------
h2 = doc.add_heading("Étape 2 : Outil 2 — Charge Interne RPE (Méthode de Foster)", level=1)
h2.paragraph_format.space_before = Pt(14)
h2.paragraph_format.space_after = Pt(6)

p2 = doc.add_paragraph(
    "La méthode de Foster (session-RPE, 2001) quantifie la charge interne globale en multipliant la durée de la séance (en minutes) par le ressenti de difficulté sur l'échelle de Borg CR-10 (1 à 10). "
    "Elle permet de suivre la charge cumulée hebdomadaire, la monotonie d'entraînement et la contrainte globale (strain)."
)

add_callout(doc, "Formules et Paramétrage dans la feuille '02_Charge_Foster_RPE'",
    "1. Structure : Date (A), Semaine (B), Joueur (C), Type de séance (D), Durée min (E), RPE (F), Charge u.a. (G), Niveau (H), Commentaire (I).\n"
    "2. Calcul de la charge de séance en G5 : =E5*F5\n"
    "3. Qualification automatique en H5 : =SI(G5>=700; \"Très Élevée\"; SI(G5>=400; \"Modérée/Élevée\"; SI(G5>0; \"Légère\"; \"Repos\")))\n"
    "4. Synthèse Hebdomadaire S1 par joueur :\n"
    "   - Charge Cumulée S1 : =SOMME.SI.ENS(G5:G40; C5:C40; \"Lucas\"; B5:B40; \"S1\")\n"
    "   - Charge Moyenne Quotidienne : =L5/7\n"
    "   - Détection Semaine Critique : =SI(L5>=2000; \"SEMAINE TRÈS CHARGÉE\"; \"Charge Normale\")")

# ---------------------------------------------------------------------------
# ÉTAPE 3
# ---------------------------------------------------------------------------
h3 = doc.add_heading("Étape 3 : Outil 3 — Journal Numérique d'Entraînement Multi-Paramètres", level=1)
h3.paragraph_format.space_before = Pt(14)
h3.paragraph_format.space_after = Pt(6)

p3 = doc.add_paragraph(
    "Dans les sports hybrides ou l'athlétisme (cas d'Antoine, 800m), le préparateur doit combiner des données de course (volume kilométrique) et des données de musculation/force (nombre d'exercices) "
    "tout en calculant la charge interne globale."
)

add_callout(doc, "Astuce Technique : Gérer les unités hétérogènes (km vs exercices)",
    "Pour calculer uniquement le kilométrage de course à pied sans additionner le nombre d'exercices de musculation, utilisez la formule SOMME.SI filtrée sur l'unité :\n"
    "=SOMME.SI(G5:G17; \"km\"; F5:F17)\n"
    "Cette formule examine la colonne Unité (G) et n'additionne dans la colonne Volume (F) que les lignes correspondant à des kilomètres de course à pied.")

# ---------------------------------------------------------------------------
# ÉTAPE 4
# ---------------------------------------------------------------------------
h4 = doc.add_heading("Étape 4 : Outil 4 — Suivi Nutritionnel & Hydratation Automatisé", level=1)
h4.paragraph_format.space_before = Pt(14)
h4.paragraph_format.space_after = Pt(6)

p4 = doc.add_paragraph(
    "Dans le cas de Sarah (triathlète), l'objectif est d'éviter la saisie manuelle fastidieuse des calories à chaque aliment. "
    "Nous construisons une table de référence nutritionnelle pour 100g, et automatisons le calcul calorique grâce à la fonction RECHERCHEV."
)

add_callout(doc, "Formule d'Automatisation RECHERCHEV",
    "Pour calculer les calories de chaque aliment consommé en fonction de la portion :\n"
    "=F5 * RECHERCHEV(E5; $M$5:$Q$16; 2; FAUX) / 100\n"
    "- E5 : Nom de l'aliment consommé.\n"
    "- $M$5:$Q$16 : Plage verrouillée de la table des aliments de référence.\n"
    "- 2 : Numéro de la colonne contenant les Kcal pour 100g.\n"
    "- FAUX : Recherche de la correspondance exacte.\n"
    "- / 100 : Conversion de la base 100g vers la quantité réelle pesée en grammes.")

# ---------------------------------------------------------------------------
# ÉTAPE 5
# ---------------------------------------------------------------------------
h5 = doc.add_heading("Étape 5 : Outil 5 — Batterie de Tests Physiques & Profil Athlétique", level=1)
h5.paragraph_format.space_before = Pt(14)
h5.paragraph_format.space_after = Pt(6)

p5 = doc.add_paragraph(
    "Les données brutes initiales mélangent les tests de vitesse (Sprint 30m) et d'explosivité (CMJ) sous un format vertical peu lisible. "
    "La première étape consiste à créer deux tableaux distincts : un tableau pour le Sprint et un tableau pour le Saut CMJ."
)

add_callout(doc, "Calculs de Progression & Particularités Sprint vs Saut",
    "Attention à la polarité des tests physiques :\n"
    "1. Pour le Sprint 30m : L'objectif est de courir plus vite (temps inférieur). La progression en secondes est donc =(Temps_Initial - Temps_Final), et le meilleur temps est =MIN(B6:E6).\n"
    "2. Pour le Saut CMJ : L'objectif est de sauter plus haut (distance supérieure). La progression en cm est =(Hauteur_Finale - Hauteur_Initiale), et le meilleur saut est =MAX(B17:E17).\n"
    "3. Pourcentage de progression : =(Progression / Valeur_Initiale) affiché avec le format personnalisé +0.0%;-0.0%;0.0%.")

# ---------------------------------------------------------------------------
# ÉTAPE 6
# ---------------------------------------------------------------------------
h6 = doc.add_heading("Étape 6 : Tableau de Bord Décisionnel Staff (Vue Direction)", level=1)
h6.paragraph_format.space_before = Pt(14)
h6.paragraph_format.space_after = Pt(6)

p6 = doc.add_paragraph(
    "Le staff technique et l'entraîneur principal n'ont pas le temps d'ouvrir les 5 feuilles de calcul détaillées. "
    "L'onglet '00_Tableau_de_Bord_Staff' regroupe en une page synthétique les indicateurs clés (KPI) et une matrice des alertes prioritaires."
)

add_callout(doc, "Check-list de Validation Finale de votre Fichier",
    "Avant de soumettre votre travail dans l'Espace Membre, vérifiez les points suivants :\n"
    "✓ Aucune cellule ne contient d'erreur (#N/A, #VALEUR!, #REF!).\n"
    "✓ Toutes les formules utilisent des références relatives et absolues appropriées ($).\n"
    "✓ Les mises en forme conditionnelles (Rouge / Orange / Vert) permettent une lecture instantanée.\n"
    "✓ Le tableau de bord synthétique met en avant les recommandations concrètes pour le staff.")

# Save document
target_local = r"docs\public\documents\Exercice-04-Tutoriel.docx"
target_drive = r"C:\Google Drive\Prépas light\HECh\Préparateur physique\Exercices\Exercice 04\Exercice-04-Tutoriel.docx"

doc.save(target_local)
doc.save(target_drive)
print(f"SUCCESS: Created {target_local} and {target_drive}")
