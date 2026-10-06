import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_exercice_02a_docx():
    doc = docx.Document()
    
    # Configuration des marges standards (2.5 cm / ~0.984 inch comme les autres exercices)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.984)
        section.bottom_margin = Inches(0.984)
        section.left_margin = Inches(0.984)
        section.right_margin = Inches(0.984)
        
    def add_p(text="", bold=False, italic=False, space_after=Pt(6)):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = 1.15
        if text:
            r = p.add_run(text)
            r.font.name = "Segoe UI"
            r.bold = bold
            r.italic = italic
        return p

    def add_bullet(text, bold_prefix="", space_after=Pt(3)):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = space_after
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.bold = True
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        return p

    # 1. Titre Principal
    p_title = add_p("Exercice 02.a — Extraction Power Query & Monitoring Nutritionnel (Natation)", bold=True, space_after=Pt(12))

    # 2. Cas professionnel
    add_p("Cas professionnel", bold=True, space_after=Pt(4))
    
    p_cas1 = doc.add_paragraph()
    p_cas1.paragraph_format.space_before = Pt(0)
    p_cas1.paragraph_format.space_after = Pt(6)
    p_cas1.paragraph_format.line_spacing = 1.15
    r = p_cas1.add_run("Vous êtes ")
    r.font.name = "Segoe UI"
    r = p_cas1.add_run("préparateur physique en charge d'un collectif de 8 nageurs de haut niveau")
    r.font.name = "Segoe UI"
    r.bold = True
    r = p_cas1.add_run(" (sprint et demi-fond). Le staff technique organise un cycle intensif pré-compétitif de 7 jours comprenant des doubles séances quotidiennes en bassin et du renforcement musculaire à sec.")
    r.font.name = "Segoe UI"

    p_cas2 = doc.add_paragraph()
    p_cas2.paragraph_format.space_before = Pt(0)
    p_cas2.paragraph_format.space_after = Pt(6)
    p_cas2.paragraph_format.line_spacing = 1.15
    r = p_cas2.add_run("Pour prévenir les risques de surentraînement, le syndrome de déficit énergétique relatif dans le sport (RED-S) et optimiser la restauration glycogénique, chaque athlète saisit quotidiennement son journal nutritionnel et ses perceptions de récupération.")
    r.font.name = "Segoe UI"

    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(0)
    p_obj.paragraph_format.space_after = Pt(12)
    p_obj.paragraph_format.line_spacing = 1.15
    r = p_obj.add_run("Objectif : ")
    r.font.name = "Segoe UI"
    r.bold = True
    r = p_obj.add_run("automatiser dans Microsoft Excel l'importation et la transformation des données nutritionnelles brutes grâce à l'outil Power Query, calculer les ratios nutritionnels relatifs au poids corporel et fournir un diagnostic décisionnel au staff d'entraînement.")
    r.font.name = "Segoe UI"

    # 3. Étape 1 : Importer et transformer les données avec Power Query
    add_p("1. Importer et transformer les données avec Power Query", bold=True, space_after=Pt(4))
    add_p("À partir du fichier fourni « Exercice 02a - Donnees brutes nutrition natation.xlsx » :", space_after=Pt(4))
    add_bullet("Dans Excel, accédez à l'onglet Données > Obtenir des données > À partir d'un fichier > À partir d'un classeur Excel.", "Importation : ")
    add_bullet("Sélectionnez la table « Export_Brut_Nutrition » et cliquez sur Transformer les données pour ouvrir l'éditeur Power Query.", "Navigateur : ")
    add_bullet("Contrôlez les types détectés par Power Query (Date, Texte, Nombre décimal pour le poids et les distances, Nombre entier pour les calories et macronutriments). Corrigez manuellement si nécessaire.", "Typage strict : ")
    add_bullet("Assurez-vous qu'aucune valeur aberrante, négative ou manquante ne corrompt le jeu de données.", "Nettoyage : ", space_after=Pt(10))

    # 4. Étape 2 : Créer des colonnes calculées relatives au poids de corps
    add_p("2. Créer des indicateurs nutritionnels personnalisés", bold=True, space_after=Pt(4))
    add_p("En préparation physique, évaluer un apport nutritionnel brut en grammes n'a pas de sens physiologique sans le rapporter au gabarit de l'athlète. Dans l'éditeur Power Query (onglet Ajouter une colonne > Colonne personnalisée), calculez :", space_after=Pt(4))
    add_bullet("[Glucides_g] / [Poids_kg] (cible recommandée en natation : 6.0 à 10.0 g/kg/j selon le volume kilométrique).", "Apport glucidique relatif [Glucides_g_kg] : ")
    add_bullet("[Proteines_g] / [Poids_kg] (cible recommandée : 1.6 à 2.2 g/kg/j pour la synthèse musculaire).", "Apport protéique relatif [Proteines_g_kg] : ")
    add_bullet("[Lipides_g] / [Poids_kg] (maintien hormonal et énergétique).", "Apport lipidique relatif [Lipides_g_kg] : ")
    add_bullet("Ajoutez une colonne conditionnelle [Alerte_Glucides] : si la charge d'entraînement [Charge_RPE_UA] >= 600 et [Glucides_g_kg] < 6.0, affichez « Alerte Déficit », sinon « Conforme ».", "Règle d'alerte conditionnelle : ", space_after=Pt(10))

    # 5. Étape 3 : Charger et organiser le classeur
    add_p("3. Chargement et modélisation sous Excel", bold=True, space_after=Pt(4))
    add_bullet("Cliquez sur Fermer et charger dans... et chargez les données sous la forme d'un Tableau Excel dans une feuille nommée « Extraction_PowerQuery ».", "Chargement : ")
    add_bullet("L'extraction doit rester dynamique : si de nouvelles lignes sont ajoutées au fichier source, un simple clic droit > Actualiser doit mettre à jour l'ensemble de vos calculs.", "Automatisation : ", space_after=Pt(10))

    # 6. Étape 4 : Analyse, interprétation et aide à la décision
    add_p("4. Interprétation professionnelle et recommandations pour le staff", bold=True, space_after=Pt(4))
    add_p("En vous appuyant sur votre tableau extrait et sur l'onglet de référence « Profils_Reperes_Staff », répondez aux 4 questions professionnelles suivantes :", space_after=Pt(4))
    add_bullet("Quels athlètes présentent un déficit glucidique critique lors des journées à haute charge (Mardi et Jeudi) ? Quelles conséquences directes ce déficit entraîne-t-il sur la puissance musculaire en bassin et la surcompensation glycogénique ?", "Question 1 — Disponibilité énergétique : ")
    add_bullet("Les cibles protéiques (1.6 à 2.2 g/kg/j) sont-elles atteintes pour les spécialistes de force et sprint (Maxime, Alexandre, Emma) ? Pourquoi le timing d'ingestion post-séance de musculation est-il déterminant ?", "Question 2 — Récupération neuromusculaire : ")
    add_bullet("Identifiez les nageurs dont l'hydratation quotidienne est insuffisante au regard du kilométrage nagé. Quel est l'impact d'une hypohydratation de 2 % sur le RPE de Foster ?", "Question 3 — Hydratation & Thermorégulation : ")
    add_bullet("Rédigez 3 préconisations opérationnelles individualisées à l'attention de l'entraîneur principal et des nageurs pour la semaine de tapering suivante.", "Question 4 — Recommandations staff : ", space_after=Pt(12))

    # 7. Modalités de remise
    add_p("Modalités de remise & Barème (sur 5 points)", bold=True, space_after=Pt(4))
    add_bullet("Déposez sur votre Espace Membre votre classeur Excel (.xlsx) contenant la requête Power Query active et vos réponses rédigées (sur une feuille dédiée ou dans un document Word d'analyse).", "Livrable attendu : ")
    add_bullet("Automatisation Power Query et colonnes calculées (2 pts) | Justesse du diagnostic physiologique (1.5 pt) | Pertinence des recommandations staff (1 pt) | Rigueur et présentation (0.5 pt).", "Barème officiel : ")

    target_path = "docs/public/documents/Exercice-02a.docx"
    doc.save(target_path)
    print(f"Exercice-02a.docx mis en page avec succès dans le style standard : {target_path}")

if __name__ == "__main__":
    create_exercice_02a_docx()
