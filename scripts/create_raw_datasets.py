import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="475569")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
thin_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
thick_bottom = Border(bottom=Side(border_style="medium", color="1E3A8A"), top=thin_side, left=thin_side, right=thin_side)

# ==============================================================================
# 1. EXERCICE 03 : DONNÉES BRUTES U18
# ==============================================================================
wb3_source = openpyxl.load_workbook(r"docs\public\documents\Exercice 03 - Production attendue.xlsx", data_only=True)
ws3_raw_src = wb3_source['Données brutes']

wb3_dest = openpyxl.Workbook()
ws3_dest = wb3_dest.active
ws3_dest.title = "Données_Brutes_U18"
ws3_dest.views.sheetView[0].showGridLines = True

ws3_dest["A1"] = "EXERCICE 03 — DONNÉES BRUTES U18 (CYCLE DE 4 SEMAINES)"
ws3_dest["A1"].font = TITLE_FONT
ws3_dest["A2"] = "Jeu de données initial à importer et nettoyer sous Microsoft Excel (Tidy Data)"
ws3_dest["A2"].font = SUBTITLE_FONT

for r in range(1, ws3_raw_src.max_row + 1):
    for c in range(1, ws3_raw_src.max_column + 1):
        val = ws3_raw_src.cell(r, c).value
        cell = ws3_dest.cell(r + 3, c, val)
        if r == 1:
            cell.fill = NAVY_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thick_bottom
        else:
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center" if c not in [2, 4] else "left")

for col in ws3_dest.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws3_dest.column_dimensions[col_letter].width = max(max_len + 3, 12)

f3_local = r"docs\public\documents\Exercice 03 - Donnees brutes U18.xlsx"
f3_drive = r"C:\Google Drive\Prépas light\HECh\Préparateur physique\Exercices\Exercice 03\Exercice 03 - Donnees brutes U18.xlsx"
wb3_dest.save(f3_local)
try:
    wb3_dest.save(f3_drive)
except Exception as e:
    print(f"Drive warning ex03: {e}")
print(f"Created: {f3_local}")

# ==============================================================================
# 2. EXERCICE 04 : DONNÉES BRUTES OUTILS NUMÉRIQUES
# ==============================================================================
wb4 = openpyxl.Workbook()
wb4.remove(wb4.active)

# Sheet 1: Hooper
ws4_h = wb4.create_sheet(title="01_Donnees_Hooper")
ws4_h.views.sheetView[0].showGridLines = True
ws4_h["A1"] = "OUTIL 1 : DONNÉES BRUTES HOOPER (U18 FOOTBALL)"
ws4_h["A1"].font = TITLE_FONT
ws4_h["A2"] = "Cotation quotidienne de 1 (Très bon) à 7 (Très mauvais)"
ws4_h["A2"].font = SUBTITLE_FONT

h_headers = ["Date", "Athlète", "Sommeil (1-7)", "Stress (1-7)", "Fatigue (1-7)", "Courbatures (1-7)"]
for ci, h in enumerate(h_headers, start=1):
    c = ws4_h.cell(4, ci, h)
    c.fill = NAVY_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = thick_bottom

raw_hooper = [
    ("2026-09-08", "Lucas", 2, 2, 3, 2),
    ("2026-09-08", "Noah", 3, 2, 3, 3),
    ("2026-09-08", "Hugo", 2, 1, 2, 2),
    ("2026-09-08", "Nathan", 3, 3, 4, 3),
    ("2026-09-09", "Lucas", 3, 2, 4, 4),
    ("2026-09-09", "Noah", 4, 3, 5, 4),
    ("2026-09-09", "Hugo", 3, 2, 3, 3),
    ("2026-09-09", "Nathan", 4, 4, 5, 5),
    ("2026-09-10", "Lucas", 4, 3, 5, 5),
    ("2026-09-10", "Noah", 5, 4, 6, 6),
    ("2026-09-10", "Hugo", 3, 2, 4, 3),
    ("2026-09-10", "Nathan", 5, 5, 6, 6),
    ("2026-09-11", "Lucas", 2, 2, 3, 3),
    ("2026-09-11", "Noah", 3, 2, 3, 3),
    ("2026-09-11", "Hugo", 2, 1, 2, 2),
    ("2026-09-11", "Nathan", 3, 2, 4, 3),
    ("2026-09-12", "Lucas", 5, 4, 6, 6),
    ("2026-09-12", "Noah", 6, 5, 7, 6),
    ("2026-09-12", "Hugo", 4, 3, 5, 4),
    ("2026-09-12", "Nathan", 6, 5, 7, 7),
    ("2026-09-13", "Lucas", 4, 3, 5, 5),
    ("2026-09-13", "Noah", 5, 4, 6, 6),
    ("2026-09-13", "Hugo", 3, 2, 4, 3),
    ("2026-09-13", "Nathan", 5, 4, 5, 6),
    ("2026-09-14", "Lucas", 2, 2, 2, 2),
    ("2026-09-14", "Noah", 2, 2, 3, 3),
    ("2026-09-14", "Hugo", 2, 1, 2, 2),
    ("2026-09-14", "Nathan", 3, 2, 3, 3),
]
for ri, r in enumerate(raw_hooper, start=5):
    for ci, v in enumerate(r, start=1):
        cell = ws4_h.cell(ri, ci, v)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center" if ci != 2 else "left")

for col in ws4_h.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws4_h.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Sheet 2: Foster RPE
ws4_r = wb4.create_sheet(title="02_Donnees_Foster_RPE")
ws4_r.views.sheetView[0].showGridLines = True
ws4_r["A1"] = "OUTIL 2 : DONNÉES BRUTES RPE (BASKETBALL)"
ws4_r["A1"].font = TITLE_FONT
ws4_r["A2"] = "Suivi de charge d'entraînement (Lucas, Noah, Hugo, Nathan)"
ws4_r["A2"].font = SUBTITLE_FONT

r_headers = ["Date", "Semaine", "Joueur", "Type de séance", "Durée (min)", "RPE (1-10)"]
for ci, h in enumerate(r_headers, start=1):
    c = ws4_r.cell(4, ci, h)
    c.fill = NAVY_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = thick_bottom

raw_rpe = [
    ("2026-09-08", "S1", "Lucas", "Technique", 75, 5),
    ("2026-09-08", "S1", "Noah", "Technique", 75, 6),
    ("2026-09-08", "S1", "Hugo", "Technique", 75, 5),
    ("2026-09-08", "S1", "Nathan", "Technique", 75, 5),
    ("2026-09-09", "S1", "Lucas", "Physique", 90, 7),
    ("2026-09-09", "S1", "Noah", "Physique", 90, 8),
    ("2026-09-09", "S1", "Hugo", "Physique", 90, 7),
    ("2026-09-09", "S1", "Nathan", "Physique", 90, 6),
    ("2026-09-10", "S1", "Lucas", "Tactique", 70, 5),
    ("2026-09-10", "S1", "Noah", "Tactique", 70, 5),
    ("2026-09-10", "S1", "Hugo", "Tactique", 70, 6),
    ("2026-09-10", "S1", "Nathan", "Tactique", 70, 5),
    ("2026-09-11", "S1", "Lucas", "Repos", 0, 0),
    ("2026-09-11", "S1", "Noah", "Repos", 0, 0),
    ("2026-09-11", "S1", "Hugo", "Repos", 0, 0),
    ("2026-09-11", "S1", "Nathan", "Repos", 0, 0),
    ("2026-09-12", "S1", "Lucas", "Match", 95, 9),
    ("2026-09-12", "S1", "Noah", "Match", 95, 9),
    ("2026-09-12", "S1", "Hugo", "Match", 95, 8),
    ("2026-09-12", "S1", "Nathan", "Match", 95, 9),
    ("2026-09-13", "S1", "Lucas", "Récupération", 45, 3),
    ("2026-09-13", "S1", "Noah", "Récupération", 45, 4),
    ("2026-09-13", "S1", "Hugo", "Récupération", 45, 3),
    ("2026-09-13", "S1", "Nathan", "Récupération", 45, 3),
    ("2026-09-14", "S1", "Lucas", "Repos", 0, 0),
    ("2026-09-14", "S1", "Noah", "Repos", 0, 0),
    ("2026-09-14", "S1", "Hugo", "Repos", 0, 0),
    ("2026-09-14", "S1", "Nathan", "Repos", 0, 0),
    ("2026-09-15", "S2", "Lucas", "Technique", 80, 6),
    ("2026-09-15", "S2", "Noah", "Technique", 80, 5),
    ("2026-09-15", "S2", "Hugo", "Technique", 80, 6),
    ("2026-09-15", "S2", "Nathan", "Technique", 80, 5),
    ("2026-09-16", "S2", "Lucas", "Physique", 85, 8),
    ("2026-09-16", "S2", "Noah", "Physique", 85, 7),
    ("2026-09-16", "S2", "Hugo", "Physique", 85, 8),
    ("2026-09-16", "S2", "Nathan", "Physique", 85, 7),
]
for ri, r in enumerate(raw_rpe, start=5):
    for ci, v in enumerate(r, start=1):
        cell = ws4_r.cell(ri, ci, v)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center" if ci in [1, 2, 5, 6] else "left")

for col in ws4_r.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws4_r.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Sheet 3: Antoine 800m
ws4_l = wb4.create_sheet(title="03_Donnees_Antoine_800m")
ws4_l.views.sheetView[0].showGridLines = True
ws4_l["A1"] = "OUTIL 3 : JOURNAL NUMÉRIQUE (ANTOINE 800M)"
ws4_l["A1"].font = TITLE_FONT
ws4_l["A2"] = "Volume, intensité, durée et RPE"
ws4_l["A2"].font = SUBTITLE_FONT

l_headers = ["Date", "Sportif", "Type de séance", "Contenu", "Durée (min)", "Volume", "Unité volume", "Intensité", "RPE", "Commentaires"]
for ci, h in enumerate(l_headers, start=1):
    c = ws4_l.cell(4, ci, h)
    c.fill = NAVY_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = thick_bottom

raw_log = [
    ("2026-09-01", "Antoine", "Endurance", "Course continue", 50, 8.0, "km", "Faible", 4, "Bonnes sensations"),
    ("2026-09-03", "Antoine", "Fractionné", "8 x 400 m", 65, 3.2, "km", "Élevée", 8, "Dernières répétitions difficiles"),
    ("2026-09-05", "Antoine", "Musculation", "Force jambes", 55, 6, "exercices", "Élevée", 7, "Fatigue musculaire"),
    ("2026-09-06", "Antoine", "Endurance", "Course continue", 60, 10.0, "km", "Modérée", 5, "Régulier"),
    ("2026-09-08", "Antoine", "Fractionné", "6 x 600 m", 70, 3.6, "km", "Très élevée", 9, "Très difficile"),
    ("2026-09-10", "Antoine", "Récupération", "Footing léger", 35, 5.0, "km", "Faible", 3, "Très bonnes sensations"),
    ("2026-09-12", "Antoine", "Endurance", "Course continue", 65, 11.0, "km", "Modérée", 5, "Bonne récupération"),
    ("2026-09-14", "Antoine", "Fractionné", "10 x 300 m", 60, 3.0, "km", "Élevée", 8, "Bonne séance"),
    ("2026-09-15", "Antoine", "Musculation", "Force + gainage", 50, 7, "exercices", "Modérée", 6, "Légère fatigue"),
    ("2026-09-17", "Antoine", "Test", "800 m", 40, 0.8, "km", "Maximale", 10, "Record personnel"),
    ("2026-09-18", "Antoine", "Récupération", "Footing léger", 30, 4.0, "km", "Faible", 2, "Très bonnes sensations"),
    ("2026-09-20", "Antoine", "Endurance", "Course continue", 70, 12.0, "km", "Modérée", 5, "Bonne séance"),
    ("2026-09-22", "Antoine", "Fractionné", "5 x 800 m", 75, 4.0, "km", "Très élevée", 9, "Fatigue importante"),
]
for ri, r in enumerate(raw_log, start=5):
    for ci, v in enumerate(r, start=1):
        cell = ws4_l.cell(ri, ci, v)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center" if ci in [1, 5, 6, 7, 8, 9] else "left")

for col in ws4_l.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws4_l.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Sheet 4: Sarah Nutrition
ws4_n = wb4.create_sheet(title="04_Donnees_Nutrition_Sarah")
ws4_n.views.sheetView[0].showGridLines = True
ws4_n["A1"] = "OUTIL 4 : SUIVI NUTRITIONNEL DE TERRAIN (SARAH)"
ws4_n["A1"].font = TITLE_FONT
ws4_n["A2"] = "Enregistrements journaliers et table de référence d'aliments"
ws4_n["A2"].font = SUBTITLE_FONT

n_headers = ["Date", "Heure", "Sportive", "Moment", "Aliment", "Quantité", "Unité", "Hydratation (ml)", "Commentaire"]
for ci, h in enumerate(n_headers, start=1):
    c = ws4_n.cell(4, ci, h)
    c.fill = NAVY_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = thick_bottom

raw_nutri = [
    ("2026-09-18", "07:00", "Sarah", "Petit-déjeuner", "Flocons d'avoine", 80, "g", 300, "Bonne satiété"),
    ("2026-09-18", "07:00", "Sarah", "Petit-déjeuner", "Banane", 120, "g", 0, "Fruit mûr"),
    ("2026-09-18", "07:00", "Sarah", "Petit-déjeuner", "Yaourt", 150, "g", 0, "Yaourt nature"),
    ("2026-09-18", "09:30", "Sarah", "Avant entraînement", "Barre céréalière", 50, "g", 500, "Facile à digérer"),
    ("2026-09-18", "10:00", "Sarah", "Pendant entraînement", "Boisson énergétique", 750, "ml", 750, "Bonne tolérance"),
    ("2026-09-18", "11:30", "Sarah", "Pendant entraînement", "Gel énergétique", 40, "g", 250, "Apport glucidique direct"),
    ("2026-09-18", "13:00", "Sarah", "Déjeuner", "Riz", 200, "g", 400, "Riz basmati cuit"),
    ("2026-09-18", "13:00", "Sarah", "Déjeuner", "Poulet", 150, "g", 0, "Blanc de poulet grillé"),
    ("2026-09-18", "13:00", "Sarah", "Déjeuner", "Légumes", 200, "g", 0, "Légumes vapeur variés"),
    ("2026-09-18", "16:30", "Sarah", "Collation", "Yaourt", 125, "g", 300, "Faim modérée"),
    ("2026-09-18", "16:30", "Sarah", "Collation", "Pomme", 150, "g", 0, "Collation saine"),
    ("2026-09-18", "19:30", "Sarah", "Dîner", "Pâtes", 200, "g", 400, "Féculents récupération"),
    ("2026-09-18", "19:30", "Sarah", "Dîner", "Saumon", 150, "g", 0, "Apport oméga 3"),
    ("2026-09-18", "19:30", "Sarah", "Dîner", "Légumes", 200, "g", 0, "Haricots et carottes"),
]
for ri, r in enumerate(raw_nutri, start=5):
    for ci, v in enumerate(r, start=1):
        cell = ws4_n.cell(ri, ci, v)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center" if ci in [1, 2, 6, 7, 8] else "left")

for col in ws4_n.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws4_n.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Sheet 5: Tests Physiques
ws4_t = wb4.create_sheet(title="05_Donnees_Tests_Physiques")
ws4_t.views.sheetView[0].showGridLines = True
ws4_t["A1"] = "OUTIL 5 : DONNÉES INITIALES NON RESTRUCTURÉES"
ws4_t["A1"].font = TITLE_FONT
ws4_t["A2"] = "Format brut vertical à transformer en tableau croisé pour le calcul des progressions"
ws4_t["A2"].font = SUBTITLE_FONT

t_headers = ["Date", "Sportif", "Test", "Résultat", "Unité", "Conditions", "Commentaires"]
for ci, h in enumerate(t_headers, start=1):
    c = ws4_t.cell(4, ci, h)
    c.fill = NAVY_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = thick_bottom

raw_tests = [
    ("2026-09-01", "Lucas", "Sprint 30 m", 4.42, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-09-01", "Noah", "Sprint 30 m", 4.51, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-09-01", "Hugo", "Sprint 30 m", 4.35, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-09-01", "Nathan", "Sprint 30 m", 4.60, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-09-01", "Jules", "Sprint 30 m", 4.47, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-09-01", "Maxime", "Sprint 30 m", 4.55, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-09-15", "Lucas", "Sprint 30 m", 4.38, "s", "Terrain extérieur sec", "Bonne récupération"),
    ("2026-09-15", "Noah", "Sprint 30 m", 4.48, "s", "Terrain extérieur sec", "Bonne récupération"),
    ("2026-09-15", "Hugo", "Sprint 30 m", 4.33, "s", "Terrain extérieur sec", "Bonne récupération"),
    ("2026-09-15", "Nathan", "Sprint 30 m", 4.55, "s", "Terrain extérieur sec", "Fatigue légère"),
    ("2026-09-15", "Jules", "Sprint 30 m", 4.44, "s", "Terrain extérieur sec", "Bonne récupération"),
    ("2026-09-15", "Maxime", "Sprint 30 m", 4.53, "s", "Terrain extérieur sec", "Bonne récupération"),
    ("2026-10-01", "Lucas", "Sprint 30 m", 4.31, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-10-01", "Noah", "Sprint 30 m", 4.45, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-10-01", "Hugo", "Sprint 30 m", 4.30, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-10-01", "Nathan", "Sprint 30 m", 4.52, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-10-01", "Jules", "Sprint 30 m", 4.40, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-10-01", "Maxime", "Sprint 30 m", 4.51, "s", "Terrain extérieur sec", "Bonnes conditions"),
    ("2026-10-15", "Lucas", "Sprint 30 m", 4.28, "s", "Terrain extérieur sec", "Très bonnes conditions"),
    ("2026-10-15", "Noah", "Sprint 30 m", 4.43, "s", "Terrain extérieur sec", "Très bonnes conditions"),
    ("2026-10-15", "Hugo", "Sprint 30 m", 4.27, "s", "Terrain extérieur sec", "Très bonnes conditions"),
    ("2026-10-15", "Nathan", "Sprint 30 m", 4.50, "s", "Terrain extérieur sec", "Très bonnes conditions"),
    ("2026-10-15", "Jules", "Sprint 30 m", 4.39, "s", "Terrain extérieur sec", "Très bonnes conditions"),
    ("2026-10-15", "Maxime", "Sprint 30 m", 4.48, "s", "Terrain extérieur sec", "Très bonnes conditions"),
    ("2026-09-01", "Lucas", "Saut vertical", 48, "cm", "Salle", "Bonnes conditions"),
    ("2026-09-01", "Noah", "Saut vertical", 45, "cm", "Salle", "Bonnes conditions"),
    ("2026-09-01", "Hugo", "Saut vertical", 53, "cm", "Salle", "Bonnes conditions"),
    ("2026-09-01", "Nathan", "Saut vertical", 42, "cm", "Salle", "Bonnes conditions"),
    ("2026-09-01", "Jules", "Saut vertical", 49, "cm", "Salle", "Bonnes conditions"),
    ("2026-09-01", "Maxime", "Saut vertical", 46, "cm", "Salle", "Bonnes conditions"),
    ("2026-09-15", "Lucas", "Saut vertical", 50, "cm", "Salle", "Bonne récupération"),
    ("2026-09-15", "Noah", "Saut vertical", 46, "cm", "Salle", "Bonne récupération"),
    ("2026-09-15", "Hugo", "Saut vertical", 54, "cm", "Salle", "Bonne récupération"),
    ("2026-09-15", "Nathan", "Saut vertical", 44, "cm", "Salle", "Fatigue légère"),
    ("2026-09-15", "Jules", "Saut vertical", 50, "cm", "Salle", "Bonne récupération"),
    ("2026-09-15", "Maxime", "Saut vertical", 47, "cm", "Salle", "Bonne récupération"),
    ("2026-10-01", "Lucas", "Saut vertical", 52, "cm", "Salle", "Bonnes conditions"),
    ("2026-10-01", "Noah", "Saut vertical", 48, "cm", "Salle", "Bonnes conditions"),
    ("2026-10-01", "Hugo", "Saut vertical", 55, "cm", "Salle", "Bonnes conditions"),
    ("2026-10-01", "Nathan", "Saut vertical", 45, "cm", "Salle", "Bonnes conditions"),
    ("2026-10-01", "Jules", "Saut vertical", 51, "cm", "Salle", "Bonnes conditions"),
    ("2026-10-01", "Maxime", "Saut vertical", 49, "cm", "Salle", "Bonnes conditions"),
    ("2026-10-15", "Lucas", "Saut vertical", 53, "cm", "Salle", "Très bonnes conditions"),
    ("2026-10-15", "Noah", "Saut vertical", 49, "cm", "Salle", "Très bonnes conditions"),
    ("2026-10-15", "Hugo", "Saut vertical", 57, "cm", "Salle", "Très bonnes conditions"),
    ("2026-10-15", "Nathan", "Saut vertical", 46, "cm", "Salle", "Très bonnes conditions"),
    ("2026-10-15", "Jules", "Saut vertical", 53, "cm", "Salle", "Très bonnes conditions"),
    ("2026-10-15", "Maxime", "Saut vertical", 50, "cm", "Salle", "Très bonnes conditions"),
]
for ri, r in enumerate(raw_tests, start=5):
    for ci, v in enumerate(r, start=1):
        cell = ws4_t.cell(ri, ci, v)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="center" if ci in [1, 4, 5] else "left")

for col in ws4_t.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws4_t.column_dimensions[col_letter].width = max(max_len + 3, 12)

f4_local = r"docs\public\documents\Exercice 04 - Donnees brutes.xlsx"
f4_drive = r"C:\Google Drive\Prépas light\HECh\Préparateur physique\Exercices\Exercice 04\Exercice 04 - Donnees brutes.xlsx"
wb4.save(f4_local)
try:
    wb4.save(f4_drive)
except Exception as e:
    print(f"Drive warning ex04: {e}")
print(f"Created: {f4_local}")

# ==============================================================================
# 3. EXERCICE 05 : DONNÉES SUIVI THOMAS
# ==============================================================================
wb5 = openpyxl.Workbook()
ws5 = wb5.active
ws5.title = "Suivi_Thomas"
ws5.views.sheetView[0].showGridLines = True

ws5["A1"] = "EXERCICE 05 — SUIVI LONGITUDINAL DE THOMAS (3 SEMAINES)"
ws5["A1"].font = TITLE_FONT
ws5["A2"] = "Étude de cas décisionnelle : analyse de convergence des charges, fatigue et tests physiques"
ws5["A2"].font = SUBTITLE_FONT

t5_headers = ["Indicateur Physiologique / Charge", "Semaine 1", "Semaine 2", "Semaine 3", "Tendance", "Alerte Préparateur"]
for ci, h in enumerate(t5_headers, start=1):
    c = ws5.cell(4, ci, h)
    c.fill = NAVY_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = thick_bottom

raw_t5 = [
    ("Charge d'entraînement hebdomadaire (u.a.)", "1 850 UA", "2 250 UA", "2 400 UA", "Forte hausse (+30%)", "Risque de surmenage"),
    ("RPE moyen des séances", "5,2 / 10", "6,4 / 10", "7,3 / 10", "Augmentation continue", "Effort perçu disproportionné"),
    ("Sommeil moyen quotidien", "8 h 00", "7 h 15", "6 h 20", "Dette de sommeil (-1h40)", "Déficit régénératif majeur"),
    ("Fatigue ressentie au réveil", "3 / 10", "5 / 10", "7 / 10", "Augmentation marquée", "Fatigue centrale accumulée"),
    ("Performance au test de sprint (30 m)", "4,82 s", "4,85 s", "4,93 s", "Régression (+0,11 s)", "Perte de vitesse neuromusculaire"),
    ("Performance au test de saut (CMJ)", "43 cm", "42 cm", "39 cm", "Chute de 4 cm (-9%)", "Atteinte de la force réactive"),
    ("Douleur musculaire ressentie", "2 / 10", "4 / 10", "6 / 10", "Aggravation", "Signal d'alarme lésionnel"),
]

for ri, r in enumerate(raw_t5, start=5):
    for ci, v in enumerate(r, start=1):
        cell = ws5.cell(ri, ci, v)
        cell.border = cell_border
        cell.alignment = Alignment(horizontal="left" if ci in [1, 5, 6] else "center")
        if ci == 6:
            cell.font = BOLD_FONT

for col in ws5.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws5.column_dimensions[col_letter].width = max(max_len + 3, 14)

f5_local = r"docs\public\documents\Exercice 05 - Donnees de suivi Thomas.xlsx"
f5_drive = r"C:\Google Drive\Prépas light\HECh\Préparateur physique\Exercices\Exercice 05\Exercice 05 - Donnees de suivi Thomas.xlsx"
wb5.save(f5_local)
try:
    wb5.save(f5_drive)
except Exception as e:
    print(f"Drive warning ex05: {e}")
print(f"Created: {f5_local}")
