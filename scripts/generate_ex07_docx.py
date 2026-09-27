import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="EFF6FF", border_hex="1E3A8A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(30, 58, 138)
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

doc = docx.Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Title
p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(2)
r_title = p_title.add_run("TUTORIEL MÉTHODOLOGIQUE DE RÉFÉRENCE")
r_title.bold = True
r_title.font.name = "Calibri"
r_title.font.size = Pt(18)
r_title.font.color.rgb = RGBColor(30, 58, 138)

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(12)
r_sub = p_sub.add_run("Exercice 07 : Mission Professionnelle Finale — Du Terrain à la Décision\nSpécialisation Préparateur Physique — HECh (Évaluation Finale 30 points / 30%)")
r_sub.font.name = "Calibri"
r_sub.font.size = Pt(12)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(100, 116, 139)

add_callout(doc, "Cadre de la Mission Finale", 
    "Vous êtes placé en situation réelle de responsable de la préparation physique d'un club de sport collectif de haut niveau. "
    "Durant 6 semaines, vous devez consolider et exploiter un jeu de données multisources (217 séances, 504 questionnaires de bien-être, tests physiques pré/post et temps de jeu en compétition). "
    "Votre mission ne consiste pas à empiler des chiffres, mais à filtrer le signal du bruit pour fournir au staff technique un tableau de bord d'aide à la décision opérationnelle.")

# Sommaire
p_toc = doc.add_paragraph()
p_toc.paragraph_format.space_before = Pt(6)
p_toc.paragraph_format.space_after = Pt(6)
r_toc = p_toc.add_run("Guide des 6 Étapes Clés pour Réussir votre Travail :")
r_toc.bold = True
r_toc.font.size = Pt(12)

steps = [
    "Étape 1 : Nettoyage et Contrôle Qualité des Données (Data Auditing)",
    "Étape 2 : Quantification des Charges et Modélisation du Risque (ACWR de Gabbett)",
    "Étape 3 : Traitement Longitudinal du Bien-Être (Sommeil, Fatigue, Douleurs)",
    "Étape 4 : Analyse Statistique Pré/Post des Tests Physiques (Semaine 1 vs Semaine 6)",
    "Étape 5 : Conception du Tableau de Bord Exécutif Staff & Préconisations Semaine 7",
    "Étape 6 : Prototypage de l'Application Mobile de Collecte & Restitution"
]
for s in steps:
    p_s = doc.add_paragraph(s, style='List Bullet')
    p_s.paragraph_format.space_before = Pt(1)
    p_s.paragraph_format.space_after = Pt(2)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# ---------------------------------------------------------------------------
# ÉTAPE 1
# ---------------------------------------------------------------------------
h1 = doc.add_heading("Étape 1 : Nettoyage et Contrôle Qualité des Données Brutes", level=1)
h1.paragraph_format.space_before = Pt(14)
h1.paragraph_format.space_after = Pt(6)

p1 = doc.add_paragraph(
    "Dans un environnement professionnel, les données de terrain ne sont jamais parfaites. "
    "Avant toute analyse statistique, vous devez auditer les feuilles RAW_Seances, RAW_Bien-etre et RAW_Tests pour identifier :"
)
p1_1 = doc.add_paragraph("• Les doublons de saisie éventuels sur une même date pour un même joueur.", style='List Bullet')
p1_2 = doc.add_paragraph("• Les valeurs aberrantes (par exemple un RPE supérieur à 10 ou une durée négative).", style='List Bullet')
p1_3 = doc.add_paragraph("• Les incohérences de statut (ex: Présence notée 0 mais durée réelle ou RPE renseigné).", style='List Bullet')
p1_4 = doc.add_paragraph("• Les cellules vides ou données manquantes sur les heures de sommeil.", style='List Bullet')

add_callout(doc, "Bonne Pratique Professionnelle",
    "Créez une feuille '01_Nettoyage_Donnees' listant clairement chaque anomalie repérée, la règle de correction appliquée "
    "(ex: suppression du doublon, imputation de la moyenne individuelle pour le sommeil manquant) et la justification méthodologique.")

# ---------------------------------------------------------------------------
# ÉTAPE 2
# ---------------------------------------------------------------------------
h2 = doc.add_heading("Étape 2 : Automatisation des Charges et Modélisation ACWR", level=1)
h2.paragraph_format.space_before = Pt(14)
h2.paragraph_format.space_after = Pt(6)

p2 = doc.add_paragraph(
    "La quantification de la charge interne repose sur la méthode de Foster : Charge (u.a.) = Durée réelle (min) × RPE (1-10). "
    "Pour anticiper les risques de blessures sans sous-entraîner les sportifs, nous intégrons le ratio ACWR (Acute:Chronic Workload Ratio de Gabbett, 2016) :"
)

add_callout(doc, "Paramétrage dans l'onglet '02_Calculs_Charges_ACWR'",
    "1. Charge de séance en I5 : =G5*H5 (Durée × RPE).\n"
    "2. Charge Aiguë (Acute Workload - 7 derniers jours) en J5 : Somme de la charge sur les 7 derniers jours pour l'athlète concerné.\n"
    "3. Charge Chronique (Chronic Workload - moyenne mobile 28 jours) en K5 : Moyenne des charges hebdomadaires sur 4 semaines.\n"
    "4. Ratio ACWR en L5 : =SI(K5>0; J5/K5; 1.0)\n"
    "5. Définition des Zones de Risque :\n"
    "   - ACWR < 0.8 : Zone de sous-charge (risque de désentraînement / vulnérabilité).\n"
    "   - 0.8 <= ACWR <= 1.3 : 'Sweet Spot' (zone optimale de progression et de protection).\n"
    "   - 1.3 < ACWR < 1.5 : Zone haute (surveillance accrue).\n"
    "   - ACWR >= 1.5 : 'Danger Zone' (sur-charge aiguë critique, risque de blessure multiplié par 2 à 4).")

# ---------------------------------------------------------------------------
# ÉTAPE 3
# ---------------------------------------------------------------------------
h3 = doc.add_heading("Étape 3 : Traitement Longitudinal du Bien-Être & Détection de Douleurs", level=1)
h3.paragraph_format.space_before = Pt(14)
h3.paragraph_format.space_after = Pt(6)

p3 = doc.add_paragraph(
    "La feuille RAW_Bien-etre compile 504 enregistrements quotidiens. Vous devez créer une synthèse dynamique par joueur et par semaine. "
    "Le croisement des scores de fatigue (1-8), de courbatures musculaires (1-8) et des heures de sommeil permet d'anticiper les baisses de performance avant l'apparition d'une lésion clinique."
)

add_callout(doc, "Alerte Rouge : Le Marqueur Douleur",
    "La colonne 'Douleur' (Oui/Non) est un commutateur binaire capital. Toute occurrence de 'Oui' associée à une note de courbatures >= 6 doit automatiquement "
    "déclencher une alerte médicale dans votre tableau de bord, impliquant un retrait immédiat des exercices à haute intensité excentrique.")

# ---------------------------------------------------------------------------
# ÉTAPE 4
# ---------------------------------------------------------------------------
h4 = doc.add_heading("Étape 4 : Analyse Statistique Pré/Post des Tests Physiques", level=1)
h4.paragraph_format.space_before = Pt(14)
h4.paragraph_format.space_after = Pt(6)

p4 = doc.add_paragraph(
    "L'onglet '05_Tests_Physiques_PrePost' compare les bilans initiaux de la Semaine 1 aux bilans finaux de la Semaine 6. "
    "Pour chaque test, respectez rigoureusement la polarité du calcul :"
)

add_callout(doc, "Formules des Deltas et Pourcentages de Progression",
    "• Sprint 10m (s) : Delta = Temps_S1 - Temps_S6 | Progression (%) = (Temps_S1 - Temps_S6) / Temps_S1 (Un delta positif est une amélioration).\n"
    "• Saut CMJ (cm) : Delta = CMJ_S6 - CMJ_S1 | Progression (%) = (CMJ_S6 - CMJ_S1) / CMJ_S1 (Un delta positif est une amélioration).\n"
    "• Test Yo-Yo IR1 (m) : Delta = Distance_S6 - Distance_S1 | Progression (%) = (Distance_S6 - Distance_S1) / Distance_S1.\n"
    "• Test Agilité 505 (s) : Delta = Temps_S1 - Temps_S6 | Progression (%) = (Temps_S1 - Temps_S6) / Temps_S1.\n"
    "• Ligne Totale : Utilisez la formule =MOYENNE() pour quantifier le bénéfice collectif moyen sur le groupe.")

# ---------------------------------------------------------------------------
# ÉTAPE 5
# ---------------------------------------------------------------------------
h5 = doc.add_heading("Étape 5 : Conception du Tableau de Bord Exécutif Staff & Décisions S7", level=1)
h5.paragraph_format.space_before = Pt(14)
h5.paragraph_format.space_after = Pt(6)

p5 = doc.add_paragraph(
    "L'onglet '00_Dashboard_Entraineur' est la vitrine de votre travail. Il doit pouvoir être lu et compris par le staff technique en moins de 60 secondes :"
)
p5_1 = doc.add_paragraph("• 5 Cartes KPI en haut d'écran (Effectif, Gains aérobie YoYo, Gains Vitesse 10m, Gains Détente CMJ, Alertes ACWR).", style='List Bullet')
p5_2 = doc.add_paragraph("• Une Matrice des Alertes Décisionnelles avec code couleur (Rouge = Critique, Orange = Vigilance, Vert = Optimal).", style='List Bullet')
p5_3 = doc.add_paragraph("• 4 Préconisations Précises pour la Semaine 7 (Semaine précédant un match à fort enjeu) : volume, décharge ciblée, sommeil, protocole d'activation J-1.", style='List Bullet')

# ---------------------------------------------------------------------------
# ÉTAPE 6
# ---------------------------------------------------------------------------
h6 = doc.add_heading("Étape 6 : Prototypage de l'Application Mobile de Terrain", level=1)
h6.paragraph_format.space_before = Pt(14)
h6.paragraph_format.space_after = Pt(6)

p6 = doc.add_paragraph(
    "La consigne vous demande de concevoir ou prototyper un outil mobile simple. Dans l'onglet '06_Prototype_App_Mobile', détaillez l'architecture suivante :"
)
p6_1 = doc.add_paragraph("• Vue Athlète : Formulaire simplifié (RPE après entraînement en 2 clics, check-in matinal Hooper).", style='List Bullet')
p6_2 = doc.add_paragraph("• Vue Préparateur : Tableau des présences, voyants de risque ACWR, alertes en temps réel avant le début de séance.", style='List Bullet')
p6_3 = doc.add_paragraph("• Connexion technique : Synchronisation directe avec le tableur (via Google Forms / AppSheet / Glide ou script API).", style='List Bullet')

add_callout(doc, "Critères d'Excellence pour la Note Maximale (30/30)",
    "✓ Rigueur mathématique absolue des formules sans rupture de référence.\n"
    "✓ Capacité à faire émerger des décisions concrètes d'entraînement plutôt que de simples graphiques décoratifs.\n"
    "✓ Prise en compte de la fatigue contextuelle (sommeil, postes de jeu, temps de jeu en match).\n"
    "✓ Clarté visuelle et ergonomie professionnelle du classeur Excel et du rapport écrit.")

# Save document
target_local = r"docs\public\documents\Exercice-07-Tutoriel.docx"
target_drive = r"C:\Google Drive\Prépas light\HECh\Préparateur physique\Exercices\Exercice 07\Exercice-07-Tutoriel.docx"

doc.save(target_local)
doc.save(target_drive)
print(f"SUCCESS: Created {target_local} and {target_drive}")
